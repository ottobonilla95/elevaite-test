import openpyxl
import os
import openai
import json
from openai import OpenAI

from fastapi import UploadFile, File
from dotenv import load_dotenv
import aspose.slides as slides



load_dotenv()
openai.api_key = os.getenv("OPENAI_API_KEY")
os.environ["OPENAI_API_KEY"] = os.getenv("OPENAI_API_KEY")


def convert_bytes_to_human_readable(size_in_bytes):
    for unit in ['B', 'KB', 'MB', 'GB', 'TB']:
        if size_in_bytes < 1024.0:
            return f"{size_in_bytes:.2f} {unit}"
        size_in_bytes /= 1024.0

async def upload_file(file: UploadFile = File(...)):
    try:

        #Get File Type
        file_type = file.content_type

        #Get File Size
        file.file.seek(0,2)
        size = file.file.tell()
        file.file.seek(0)
        file_size = convert_bytes_to_human_readable(size)

        #Get File Name
        file_name = file.filename

        #Save Excel file to local drive
        os.makedirs("data/Excel", exist_ok=True)
        with open(f"data/Excel/{file.filename}", "wb") as f:
            f.write(file.file.read())

        #Excel/Workbook - get list of sheets
        sheets = []
        sheets_count = 0
        if file_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or file_type == "application/vnd.ms-excel":
            file_path = f"data/Excel/{file.filename}"
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            sheets = workbook.sheetnames
            sheets_count = len(sheets)
        return {"response": "Success!", "file_size": file_size, "file_name": file_name, "file_type" : file_type, "sheet": sheets, "sheets_count": sheets_count}
    except Exception as e:
        return {"response": "Error!", "error_message": e}


def api_openai(content: str):
    client = OpenAI()
    chat_completion = client.chat.completions.create(
        messages=[
            {"role":"system", "content":"You are a helpful assistant."},
            {"role":"user","content":content}

        ],
        model="gpt-4",
    )
    return(chat_completion.choices[0].message.content)



